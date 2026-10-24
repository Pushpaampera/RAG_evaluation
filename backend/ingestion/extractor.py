from io import BytesIO
from pathlib import Path

import pdfplumber
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_document(file_bytes: bytes, filename: str) -> dict:
    """
    Extract text from an uploaded PDF or DOCX document.

    Parameters:
        file_bytes: Uploaded document content in bytes.
        filename: Original document filename.

    Returns:
        Dictionary containing extracted text and document metadata.

    Raises:
        ValueError: If the file is empty, unsupported, or contains
                    no extractable text.
    """

    
    # 1. VALIDATE INPUT
    

    if not file_bytes:
        raise ValueError("The uploaded document is empty.")

    if not filename or not filename.strip():
        raise ValueError("A valid filename is required.")

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. Please upload a PDF or DOCX file."
        )

    
    # 2. EXTRACT PDF
    
    if extension == ".pdf":
        extracted_text = []
        page_count = 0

        try:
            with pdfplumber.open(BytesIO(file_bytes)) as pdf:
                page_count = len(pdf.pages)

                if page_count == 0:
                    raise ValueError(
                        "The PDF document does not contain any pages."
                    )

                for page_number, page in enumerate(
                    pdf.pages,
                    start=1,
                ):
                    page_text = page.extract_text()

                    if page_text and page_text.strip():
                        extracted_text.append(
                            f"## Page {page_number}\n\n"
                            f"{page_text.strip()}"
                        )

        except ValueError:
            raise

        except Exception as error:
            raise ValueError(
                "The PDF could not be processed. "
                "Please check that the file is a valid PDF."
            ) from error

        text = "\n\n".join(extracted_text)

        if not text.strip():
            raise ValueError(
                "No extractable text was found in the PDF. "
                "The document may be scanned or image-based."
            )

        return {
            "filename": filename,
            "file_type": "pdf",
            "page_count": page_count,
            "text": text,
        }

   
    # 3. EXTRACT DOCX
    

    if extension == ".docx":
        extracted_text = []

        try:
            document = Document(BytesIO(file_bytes))

            # Extract paragraphs while preserving headings.
            for paragraph in document.paragraphs:
                paragraph_text = paragraph.text.strip()

                if not paragraph_text:
                    continue

                style_name = (
                    paragraph.style.name
                    if paragraph.style
                    else ""
                )

                if style_name.startswith("Heading"):
                    extracted_text.append(
                        f"## {paragraph_text}"
                    )
                else:
                    extracted_text.append(paragraph_text)

            # Extract table content.
            for table in document.tables:
                for row in table.rows:
                    cells = [
                        cell.text.strip().replace("\n", " ")
                        for cell in row.cells
                    ]

                    if any(cells):
                        extracted_text.append(
                            " | ".join(cells)
                        )

        except Exception as error:
            raise ValueError(
                "The DOCX document could not be processed. "
                "Please check that the file is a valid DOCX."
            ) from error

        text = "\n\n".join(extracted_text)

        if not text.strip():
            raise ValueError(
                "No extractable text was found in the DOCX document."
            )

        return {
            "filename": filename,
            "file_type": "docx",
            "page_count": None,
            "text": text,
        }

    raise ValueError("Unsupported document format.")